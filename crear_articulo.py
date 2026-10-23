from salidas import archivo, FECHA
from pathlib import Path
from autoria import NOMBRE, FIRMA, AUTOR
from analisis_evolucion import secciones_evolucion
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parent
out=ROOT/'informe'; out.mkdir(exist_ok=True)
summary=json.loads((ROOT/'resultados'/archivo('resumen_economico.json')).read_text(encoding='utf8'))
# Todas las cifras y descripciones se derivan del corte de esta ejecución.
def f(value,digits=2): return f'{value:.{digits}f}'.replace('-', '−').replace('.', ',')
n=summary['n']; first=summary['inicio']; last=summary['final']; ext=summary['extremos']
ipc_mode=summary.get('ipc_ajuste','original')
period_text='; '.join(f"{p['periodo']}: {f(p['correlacion'],3)} ({p['n']} meses, {p['desde']} a {p['hasta']})" for p in summary['subperiodos'])
coverage=__import__('pandas').read_csv(ROOT/'resultados'/archivo('cobertura.csv'))
missing=__import__('pandas').read_csv(ROOT/'resultados'/archivo('meses_excluidos.csv'))
manifest_path=ROOT/'fuentes'/'actualizacion.json'
manifest=json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else {}
source_dates=json.loads((ROOT/'fuentes'/'imacec_metadata.json').read_text(encoding='utf-8'))
sections=[
('title',f"Inflación, empleo y actividad en Chile: {summary['desde']} a {summary['hasta']}"),
('subtitle',f"{n} meses comunes | Edición {FECHA} | IPC {ipc_mode}"),
('p',FIRMA),
('p',AUTOR['nota']),
('p',f"La correlación contemporánea entre inflación anual y desocupación es {f(summary['correlacion'],3)} en el período {summary['desde']} a {summary['hasta']}. Es una asociación descriptiva: no identifica causalidad ni permite recomendar una tasa de política monetaria."),
('p','El eje horizontal representa la desocupación ENE y el vertical la inflación anual. El área de las burbujas es proporcional al valor absoluto de la variación anual del IR real. El color muestra la variación interanual del promedio móvil de tres niveles originales del IMACEC: naranja para tasas negativas, claro para cero y azul para positivas. Las escalas son fijas dentro de cada edición.'),
('image',''),
('caption',f"Período común {summary['desde']}–{summary['hasta']}. ENE e IMACEC promedio 3m alineados al mes central. Bordes segmentados: ventana marzo de 2020–agosto de 2023. Referencias históricas: meta de inflación 3% y punto medio NAIRU 8,25%."),
*secciones_evolucion(summary),
('h','Cobertura y fechas de las fuentes'),
('p',f"El cruce conserva {n} meses entre {summary['desde']} y {summary['hasta']}. Se excluyen {len(missing)} meses de la unión por falta de al menos una variable necesaria. El último punto es el último mes común completo; puede terminar antes que una fuente individual. No se rellenan faltantes."),
('p','ENE se asigna al mes central de su trimestre móvil. El promedio centrado del IMACEC requiere el mes siguiente: el análisis es retrospectivo y no reconstruye la información conocida en tiempo real. Los trimestres móviles y las tasas anuales comparten observaciones.'),
]
for row in coverage.itertuples(index=False):
    sections.append(('p',f"{row.serie}: {row.desde} a {row.hasta}; {row.filas} filas; {row.faltantes} valores faltantes en la variable auditada."))
if manifest:
    sections.append(('p',f"Última descarga conjunta validada de estos archivos: {manifest['descargado_utc']}. El registro de ejecución identifica si esta edición descargó fuentes o usó explícitamente el corte local. Hashes exactos de los CSV: archivo fuentes_sha256.json de esta edición."))
sections.extend([
('h','Cierres anuales y fórmulas'),
('p','Color: gₜ = 100 × (Īₜ / Īₜ₋₁₂ − 1), donde Īₜ = (Iₜ₋₁ + Iₜ + Iₜ₊₁) / 3. Se comparan promedios de niveles originales, no promedios de tasas. Panel: IMACEC desestacionalizado e IPC en variaciones a 12 meses.'),
('p','En diciembre, el rombo muestra el crecimiento del promedio anual del IMACEC original y el cuadrado la inflación diciembre/diciembre. El tamaño y color se normalizan con las tasas a 12 meses de ambas series y un máximo absoluto común fijo en toda la muestra. Solo hay puntos de cierre para diciembres incluidos; los años parciales no se presentan como años completos.'),
])
closings=[r for m,r in summary['observaciones'].items() if m.endswith('-12')]
if closings:
    closing=closings[-1]
    sections.append(('p',f"Último diciembre incluido: {closing['mes']}. IMACEC original, crecimiento del promedio anual: {f(closing['imacec_acumulado_anual'])}%; IMACEC desestacionalizado, variación a 12 meses: {f(closing['imacec_sa_anual'])}%; IPC de cierre: {f(closing['ipc_anual'])}%."))
sections.extend([
('h','Alcance de las referencias'),
('p','La meta del 3% corresponde al horizonte de dos años. La referencia NAIRU 8,25% es el punto medio propio del rango BCCh 8,0–8,5% para 2024-T3, publicado en diciembre de 2024. Se mantiene como referencia histórica y no como estimación actual ni trayectoria anual. Sus distancias son ilustrativas; no son brechas oficiales.'),
('h','Fuentes y reproducibilidad'),
('source','INE: ENE, IPC histórico empalmado e IR real empalmado. URL exactas de descarga en extractors/ene_extractor.py, extractors/ipc_extractor.py y extractors/ir_extractor.py. https://www.ine.gob.cl/'),
('source','BCCh: IMACEC original y desestacionalizado, base 2018=100. Series F032.IMC.IND.Z.Z.EP18.Z.Z.0.M y F032.IMC.IND.Z.Z.EP18.Z.Z.1.M. Las tasas calculadas con índices publicados a un decimal son aproximadas. https://www.bcentral.cl/areas/estadisticas/imacec'),
('source','Referencia histórica NAIRU: minutas citadas en IPoM diciembre de 2024, Holguras en el mercado laboral. https://www.bcentral.cl/documents/33528/6735463/Minutas%2Bcitadas%2Ben%2Bel%2BIPoM%2Bdiciembre%2B2024.pdf/d24985ae-cb5e-2f03-4499-ecfad3d86ade'),
('source','Código, cobertura, meses excluidos y hashes permiten reproducir cada edición. https://github.com/alejandrohenriquezr/curva_phillips'),
])
for name in ('imacec_original','imacec_sa'):
    item=source_dates[name]
    sections.append(('source',f"{name}: actualización informada por BDE {item.get('actualizacion_bcch')}; descarga {source_dates.get('descargado_utc')}; {item['url']}"))

# La variante SA se identifica y no se atribuye al INE.
if ipc_mode=='sa':
    meta=json.loads((ROOT/'resultados'/archivo('ipc_x13_metadatos.json')).read_text(encoding='utf8'))
    import csv
    with (ROOT/'resultados'/archivo('ipc_comparacion.csv')).open(encoding='utf-8-sig') as source:
        comparison=list(csv.DictReader(source))
    latest=comparison[-1]
    replacements={
        'Análisis para LinkedIn':'Análisis con IPC desestacionalizado experimental',
        'inflación anual en el vertical':'inflación anual del IPC ajustado X-13/SEATS en el vertical',
        'usando la tasa publicada.':'usando el índice IPC ajustado X-13/SEATS.',
        'se usa la tasa anual publicada de diciembre para evitar diferencias de redondeo.':'se calcula sobre el índice ajustado de diciembre contra diciembre.',
    }
    sections=[(kind, __import__('functools').reduce(lambda t,p:t.replace(*p),replacements.items(),text)) for kind,text in sections]
    ix=next(i for i,(kind,text) in enumerate(sections) if kind=='h' and text=='Fuentes y reproducibilidad')
    sections[ix:ix]=[
        ('h','Ajuste estacional experimental del IPC'),
        ('p',f"Esta edición utiliza X-13ARIMA-SEATS con descomposición SEATS sobre los {meta['n']} niveles empalmados mensuales de {meta['desde']} a {meta['hasta']}. El modelo regARIMA seleccionado fue {meta['modelo_regarima']}, con logaritmos y detección automática de valores atípicos AO, LS y TC. El archivo de especificación y el informe completo del programa acompañan los resultados [9]."),
        ('p','El ajuste se aplica al nivel del índice, no a las tasas publicadas. Si Pₜᴬ es el índice desestacionalizado, la variación mensual es 100 × (Pₜᴬ / Pₜ₋₁ᴬ − 1) y la de 12 meses es 100 × (Pₜᴬ / Pₜ₋₁₂ᴬ − 1). La medida de diciembre usa diciembre contra diciembre ajustados; no es la inflación acumulada oficial. Las tasas originales publicadas permanecen en el CSV comparativo.'),
        ('p',f"En {latest['fecha'][:7]}, último IPC disponible, la variación mensual publicada es {f(float(latest['ipc_mensual_original']))}% y la ajustada {f(float(latest['ipc_mensual_sa']))}%. A 12 meses son {f(float(latest['ipc_anual_original']))}% y {f(float(latest['ipc_anual_sa']))}%, respectivamente. El gráfico de Phillips termina en {last['mes']}, por la intersección con empleo y remuneraciones."),
        ('p','Las diferencias a 12 meses pueden ser pequeñas porque se compara el mismo mes del año; el ajuste suele ser más informativo al comparar meses consecutivos. Parte de la diferencia frente a la tasa publicada también procede del redondeo del índice empalmado, por lo que se exporta además su variación original recalculada.'),
        ('p','El diagnóstico advierte un pico de días de negociación en el espectro de los residuos. No se incorporaron regresores de días hábiles ni feriados específicos de Chile. La estimación es exploratoria: requiere evaluar efectos calendario, estabilidad y revisiones antes de un uso oficial. La eliminación de estacionalidad no equivale a eliminar todos los efectos de calendario.'),
        ('p','Se utiliza la muestra completa disponible, incluida información posterior a los puntos históricos. Es un ejercicio retrospectivo, sujeto a revisión al incorporar observaciones; no reconstruye lo que se conocía en tiempo real. El IR real y la ENE mantienen su tratamiento anterior. La referencia de inflación de 3% conserva su sentido oficial y la distancia calculada con IPC ajustado es solo ilustrativa.'),
        ('break','')]
    sections.append(('source','[9] U.S. Census Bureau. X-13ARIMA-SEATS v1.1 build 62 y manual de referencia. https://www.census.gov/data/software/x13as.X-13ARIMA-SEATS.html'))

doc=Document()
s=doc.sections[0]; s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.8); s.bottom_margin=Cm(1.8); s.left_margin=Cm(2); s.right_margin=Cm(2)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.10
for name,size in [('Title',25),('Subtitle',11),('Heading 1',14)]:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string('000000')
    st.paragraph_format.space_after=Pt(8)
    if name=='Heading 1': st.paragraph_format.space_before=Pt(9)
md=[]
for kind,text in sections:
    if kind=='title': doc.add_paragraph(text,'Title'); md.append('# '+text)
    elif kind=='subtitle': doc.add_paragraph(text,'Subtitle'); md.append(text)
    elif kind=='h': doc.add_paragraph(text,'Heading 1'); md.append('## '+text)
    elif kind=='break': doc.add_page_break()
    elif kind=='image':
        pic=doc.add_picture(str(ROOT/'resultados'/archivo('phillips_estatico.png')),width=Cm(17))
        pic._inline.docPr.set('descr',f"Inflación y desocupación de Chile: {summary['desde']} a {summary['hasta']}; IR real, IMACEC y panel temporal.")
        md.append(f'![Gráfico de Phillips](../resultados/{archivo("phillips_estatico.png")})')
    else:
        p=doc.add_paragraph(text)
        if kind in ('caption','source'):
            for run in p.runs: run.font.size=Pt(8.5)
            p.paragraph_format.space_after=Pt(5)
        md.append(text)
for tree in (doc.element, doc.styles.element):
    for border in list(tree.xpath('//w:pBdr')):
        border.getparent().remove(border)
doc.core_properties.title=sections[0][1]
doc.core_properties.subject='Análisis descriptivo de inflación, desocupación y remuneraciones reales en Chile'
doc.core_properties.author=NOMBRE
doc.core_properties.last_modified_by=NOMBRE
doc.save(out/archivo('Articulo_LinkedIn_Curva_Phillips.docx'))
(out/archivo('articulo_linkedin.md')).write_text('\n\n'.join(md)+'\n',encoding='utf8')
print('Artículo creado')


